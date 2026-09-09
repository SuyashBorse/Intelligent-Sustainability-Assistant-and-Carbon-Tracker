import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('Project Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. Project Title & Team Introduction
    doc.add_heading('1. Project Title & Team Introduction', level=1)
    doc.add_paragraph('Project Title: Intelligent Sustainability Assistant and Carbon Tracker', style='List Bullet')
    doc.add_paragraph('Team Introduction: [Insert Team Member Names and Roles Here]', style='List Bullet')
    
    # 2. Domain Overview
    doc.add_heading('2. Domain Overview', level=1)
    doc.add_paragraph('Domain: Environmental Sustainability, Carbon Footprint Tracking, Artificial Intelligence, and Gamification.')
    doc.add_paragraph('Overview: The domain focuses on empowering individuals to monitor and mitigate their environmental impact. By tracking daily activities across transportation, energy usage, food consumption, water usage, and shopping, users gain insights into their carbon footprint (CO2e). Integrating Artificial Intelligence offers personalized decarbonization strategies, while gamification enhances user engagement through challenges, streaks, and badges.')
    
    # 3. Problem Statement
    doc.add_heading('3. Problem Statement', level=1)
    doc.add_paragraph('Despite growing environmental awareness, individuals lack accessible, cohesive tools to accurately measure, understand, and reduce their personal carbon footprints. Existing solutions are often fragmented, overly complex, or lack actionable guidance. There is a need for an integrated platform that not only calculates real-time emissions using standardized factors but also provides personalized, AI-driven recommendations and motivational gamification to encourage long-term sustainable habits.')
    
    # 4. Literature Survey
    doc.add_heading('4. Literature Survey', level=1)
    doc.add_paragraph('Carbon Footprint Tracking Applications: Studies emphasize the importance of continuous monitoring to change consumer behavior. Traditional apps use manual entry but often face low retention rates.', style='List Bullet')
    doc.add_paragraph('AI in Sustainability: Recent research highlights the use of Large Language Models (LLMs) to provide context-aware, actionable advice. LLMs can analyze user data to suggest personalized, high-impact emission reduction strategies.', style='List Bullet')
    doc.add_paragraph('Gamification for Environmental Behavior: Gamification elements (points, badges, leaderboards) have been shown to significantly increase user motivation and adherence to pro-environmental behaviors by transforming routine tracking into an engaging experience.', style='List Bullet')
    doc.add_paragraph('Emission Calculation Standards: Utilizing standard frameworks like the GHG Protocol and IPCC emission factors ensures the scientific validity and comparability of footprint estimates.', style='List Bullet')
    
    # 5. Comparative Analysis
    doc.add_heading('5. Comparative Analysis', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Feature'
    hdr_cells[1].text = 'Traditional Carbon Trackers'
    hdr_cells[2].text = 'Proposed System'
    
    comparisons = [
        ('Data Input', 'Manual, often tedious', 'Streamlined, categorized logging'),
        ('Recommendations', 'Generic tips', 'Personalized, AI-generated (Gemini) insights'),
        ('Engagement', 'Low retention', 'High retention via Streaks, Badges, Challenges'),
        ('Real-time Analytics', 'Basic charts', 'Interactive Chart.js visualizations'),
        ('Goal Setting', 'Basic targets', 'Granular tracking with progress updates')
    ]
    
    for feature, traditional, proposed in comparisons:
        row_cells = table.add_row().cells
        row_cells[0].text = feature
        row_cells[1].text = traditional
        row_cells[2].text = proposed
        
    # 6. Proposed Methodology
    doc.add_heading('6. Proposed Methodology', level=1)
    doc.add_paragraph('The project adopts a modern web architecture focusing on real-time responsiveness and data privacy.')
    doc.add_paragraph('Frontend: Built with HTML5, CSS3, and Vanilla JavaScript (ES modules). It utilizes Chart.js for interactive analytics and visualizations.', style='List Bullet')
    doc.add_paragraph('Backend & Database: Supabase (PostgreSQL) is used for secure user authentication, data persistence, and Row Level Security (RLS) to ensure users only access their own data.', style='List Bullet')
    doc.add_paragraph('AI Integration: The Google Gemini API serves as the "AI Coach," analyzing aggregated user data to provide structured, personalized sustainability recommendations.', style='List Bullet')
    doc.add_paragraph('Calculation Engine: A dedicated module calculates real-time CO2e based on standardized emission factors (e.g., IPCC) across five categories: Transportation, Energy, Food, Water, and Shopping.', style='List Bullet')
    
    # 7. Software and Hardware Requirements
    doc.add_heading('7. Software and Hardware Requirements', level=1)
    doc.add_heading('Hardware Requirements:', level=2)
    doc.add_paragraph('Processor: Intel Core i3 / AMD Ryzen 3 or higher.', style='List Bullet')
    doc.add_paragraph('RAM: 4 GB minimum (8 GB recommended).', style='List Bullet')
    doc.add_paragraph('Storage: 20 GB available space.', style='List Bullet')
    doc.add_paragraph('Internet connection.', style='List Bullet')
    
    doc.add_heading('Software Requirements:', level=2)
    doc.add_paragraph('Operating System: Windows 10/11, macOS, or Linux.', style='List Bullet')
    doc.add_paragraph('Web Browser: Modern browser (Google Chrome, Mozilla Firefox, Safari).', style='List Bullet')
    doc.add_paragraph('Development Environment: Visual Studio Code, Node.js (for Vite dev server).', style='List Bullet')
    doc.add_paragraph('Backend Services: Supabase account.', style='List Bullet')
    doc.add_paragraph('API Services: Google Gemini API Key.', style='List Bullet')
    
    # 8. Timeline / Gantt Chart
    doc.add_heading('8. Timeline / Gantt Chart', level=1)
    doc.add_paragraph('Phase 1: Requirement Analysis & Design (Weeks 1-2): PRD creation, architecture design, UI/UX mockups, and database schema definition.', style='List Bullet')
    doc.add_paragraph('Phase 2: Frontend Development & Core Logic (Weeks 3-5): HTML/CSS setup, implementation of the calculation engine, and activity logging interface.', style='List Bullet')
    doc.add_paragraph('Phase 3: Backend & AI Integration (Weeks 6-7): Supabase authentication, database connection, and integrating the Google Gemini API for personalized coaching.', style='List Bullet')
    doc.add_paragraph('Phase 4: Gamification & Analytics (Weeks 8-9): Developing point systems, badges, streaks, and integrating Chart.js for data visualization.', style='List Bullet')
    doc.add_paragraph('Phase 5: Testing & Deployment (Weeks 10-11): Manual and automated testing, bug fixing, and final deployment to a hosting platform.', style='List Bullet')
    
    # 9. References
    doc.add_heading('9. References', level=1)
    doc.add_paragraph('Greenhouse Gas Protocol (GHG Protocol) Standards.', style='List Number')
    doc.add_paragraph('Intergovernmental Panel on Climate Change (IPCC) Emission Factor Database.', style='List Number')
    doc.add_paragraph('Google Gemini API Documentation for prompt engineering and structured JSON outputs.', style='List Number')
    doc.add_paragraph('Chart.js Documentation for data visualization.', style='List Number')
    doc.add_paragraph('Supabase Documentation for PostgreSQL and Row Level Security implementation.', style='List Number')
    
    # Save the document
    file_path = "Project_Report.docx"
    doc.save(file_path)
    print(f"Document saved to {file_path}")

if __name__ == "__main__":
    create_report()
